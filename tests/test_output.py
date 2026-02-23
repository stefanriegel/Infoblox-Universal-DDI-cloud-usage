"""
Tests for output pipeline: XLS report, estimator CSV, and proof manifest.

Verifies professional formatting, correct row/column counts, token calculations,
and cryptographic proof integrity across all output formats.
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

import csv
import hashlib
import json
import os
import tempfile

import openpyxl

from cloud_usage.schema.resource import CloudResource
from cloud_usage.output.xlsx_report import write_xlsx_report
from cloud_usage.output.estimator_csv import write_estimator_csv
from cloud_usage.output.proof_manifest import write_proof_manifest


def _make_resource(
    resource_id: str = "arn:aws:ec2:us-east-1:111111111111:instance/i-abc",
    resource_type: str = "ec2-instance",
    provider: str = "aws",
    account_id: str = "111111111111",
    region: str = "us-east-1",
    name: str = "test-instance",
    ip_addresses: list[str] | None = None,
    tags: dict[str, str] | None = None,
    counted: bool = True,
    category: str | None = "asset",
    skip_reason: str | None = None,
) -> CloudResource:
    """Helper to create a CloudResource for output tests."""
    return CloudResource(
        resource_id=resource_id,
        resource_type=resource_type,
        provider=provider,
        account_id=account_id,
        region=region,
        name=name,
        ip_addresses=ip_addresses or [],
        tags=tags or {},
        discovered_at="2026-02-23T10:00:00",
        counted=counted,
        category=category,
        skip_reason=skip_reason,
    )


def _sample_resources() -> list[CloudResource]:
    """Create a realistic set of resources for testing."""
    return [
        _make_resource(
            resource_id="arn:aws:ec2:us-east-1:111111111111:vpc/vpc-abc",
            resource_type="vpc",
            name="main-vpc",
            counted=True,
            category="ddi",
        ),
        _make_resource(
            resource_id="arn:aws:ec2:us-east-1:111111111111:subnet/subnet-abc",
            resource_type="subnet",
            name="public-subnet",
            counted=True,
            category="ddi",
        ),
        _make_resource(
            resource_id="arn:aws:ec2:us-east-1:111111111111:instance/i-abc",
            resource_type="ec2-instance",
            name="web-server",
            ip_addresses=["10.0.1.5", "54.23.100.50"],
            tags={"Environment": "production", "Team": "platform"},
            counted=True,
            category="asset",
        ),
        _make_resource(
            resource_id="arn:aws:ec2:us-east-1:111111111111:volume/vol-abc",
            resource_type="ebs-volume",
            name="data-volume",
            counted=False,
            category=None,
            skip_reason="token-free: EBS Volume",
        ),
        _make_resource(
            resource_id="arn:aws:ec2:us-west-2:222222222222:vpc/vpc-xyz",
            resource_type="vpc",
            account_id="222222222222",
            region="us-west-2",
            name="secondary-vpc",
            counted=True,
            category="ddi",
        ),
    ]


def _sample_account_summaries() -> dict[str, dict]:
    """Create account summaries matching _sample_resources."""
    return {
        "111111111111": {
            "ddi_count": 2,
            "ip_count": 2,
            "asset_count": 1,
            "ddi_tokens": 1,
            "ip_tokens": 1,
            "asset_tokens": 1,
            "total_tokens": 3,
        },
        "222222222222": {
            "ddi_count": 1,
            "ip_count": 0,
            "asset_count": 0,
            "ddi_tokens": 1,
            "ip_tokens": 0,
            "asset_tokens": 0,
            "total_tokens": 1,
        },
    }


def _sample_errors() -> list[dict[str, str]]:
    """Create sample discovery errors."""
    return [
        {
            "account": "333333333333",
            "region": "eu-west-1",
            "resource_type": "ec2-instance",
            "error": "AccessDenied: not authorized",
            "suggestion": "Check IAM role permissions",
        },
    ]


class TestXlsxReportGeneration:
    """Test XLS report file creation and basic structure."""

    def test_write_xlsx_creates_file(self, tmp_path):
        """Report file is created and non-empty."""
        filepath = str(tmp_path / "test_report.xlsx")
        result = write_xlsx_report(
            filepath, _sample_resources(), _sample_account_summaries(),
            _sample_errors(), "aws",
        )
        assert result == filepath
        assert os.path.exists(filepath)
        assert os.path.getsize(filepath) > 0

    def test_write_xlsx_returns_filepath(self, tmp_path):
        """write_xlsx_report returns the filepath written."""
        filepath = str(tmp_path / "test_report.xlsx")
        result = write_xlsx_report(
            filepath, [], {}, [], "aws",
        )
        assert result == filepath


class TestXlsxDetailSheet:
    """Test the Detail sheet content and structure."""

    def test_detail_sheet_exists(self, tmp_path):
        """Detail sheet is present in the workbook."""
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, _sample_resources(),
                          _sample_account_summaries(), [], "aws")
        wb = openpyxl.load_workbook(filepath)
        assert "Detail" in wb.sheetnames
        wb.close()

    def test_detail_row_count(self, tmp_path):
        """Detail sheet has header + one row per resource."""
        resources = _sample_resources()
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, resources,
                          _sample_account_summaries(), [], "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Detail"]
        # Count non-empty rows
        row_count = 0
        for row in ws.iter_rows(min_row=1, max_col=1):
            if row[0].value is not None:
                row_count += 1
        assert row_count == len(resources) + 1  # header + resources
        wb.close()

    def test_detail_header_columns(self, tmp_path):
        """Detail sheet headers match expected columns."""
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, _sample_resources(),
                          _sample_account_summaries(), [], "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Detail"]
        headers = [cell.value for cell in ws[1]]
        expected = [
            "Resource ID", "Type", "Account", "Region", "Name",
            "IP Addresses", "IP Count", "Counted", "Category",
            "Skip Reason", "Tags",
        ]
        assert headers == expected
        wb.close()

    def test_detail_counted_yes_no(self, tmp_path):
        """Counted column shows 'Yes' or 'No' correctly."""
        resources = [
            _make_resource(counted=True, category="ddi"),
            _make_resource(
                resource_id="vol-1", counted=False, category=None,
                skip_reason="token-free",
            ),
        ]
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, resources, {}, [], "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Detail"]
        # Counted is column H (index 8, 1-based)
        assert ws.cell(row=2, column=8).value == "Yes"
        assert ws.cell(row=3, column=8).value == "No"
        wb.close()

    def test_detail_ip_addresses_and_count(self, tmp_path):
        """IP Addresses are comma-separated and IP Count is correct."""
        resources = [
            _make_resource(ip_addresses=["10.0.1.5", "54.23.100.50"]),
        ]
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, resources, {}, [], "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Detail"]
        # IP Addresses is column F (6), IP Count is column G (7)
        assert ws.cell(row=2, column=6).value == "10.0.1.5, 54.23.100.50"
        assert ws.cell(row=2, column=7).value == 2
        wb.close()

    def test_detail_tags_formatted(self, tmp_path):
        """Tags are formatted as key=val; key2=val2."""
        resources = [
            _make_resource(tags={"Env": "prod", "Team": "ops"}),
        ]
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, resources, {}, [], "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Detail"]
        # Tags is column K (11)
        tags_val = ws.cell(row=2, column=11).value
        assert "Env=prod" in tags_val
        assert "Team=ops" in tags_val
        wb.close()

    def test_detail_category_uppercase(self, tmp_path):
        """Category displays in uppercase or '-' for None."""
        resources = [
            _make_resource(counted=True, category="ddi"),
            _make_resource(
                resource_id="vol-1", counted=False, category=None,
                skip_reason="token-free",
            ),
        ]
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, resources, {}, [], "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Detail"]
        # Category is column I (9)
        assert ws.cell(row=2, column=9).value == "DDI"
        assert ws.cell(row=3, column=9).value == "-"
        wb.close()

    def test_detail_skip_reason(self, tmp_path):
        """Skip reason appears for excluded resources, empty for counted."""
        resources = [
            _make_resource(counted=True, category="ddi", skip_reason=None),
            _make_resource(
                resource_id="vol-1", counted=False, category=None,
                skip_reason="token-free: EBS Volume",
            ),
        ]
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, resources, {}, [], "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Detail"]
        # Skip Reason is column J (10)
        # openpyxl reads empty strings as None
        assert ws.cell(row=2, column=10).value in ("", None)
        assert ws.cell(row=3, column=10).value == "token-free: EBS Volume"
        wb.close()


class TestXlsxSummarySheet:
    """Test the Summary sheet content and structure."""

    def test_summary_sheet_exists(self, tmp_path):
        """Summary sheet is present in the workbook."""
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, [], _sample_account_summaries(), [], "aws")
        wb = openpyxl.load_workbook(filepath)
        assert "Summary" in wb.sheetnames
        wb.close()

    def test_summary_account_count(self, tmp_path):
        """Summary has one row per account plus header and total row."""
        summaries = _sample_account_summaries()
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, [], summaries, [], "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Summary"]
        row_count = 0
        for row in ws.iter_rows(min_row=1, max_col=1):
            if row[0].value is not None:
                row_count += 1
        # header + accounts + total row
        assert row_count == 1 + len(summaries) + 1
        wb.close()

    def test_summary_total_row(self, tmp_path):
        """Provider total row shows correct sums."""
        summaries = _sample_account_summaries()
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, [], summaries, [], "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Summary"]
        # Total row is at row = 1 (header) + len(summaries) + 1
        total_row = len(summaries) + 2
        assert ws.cell(row=total_row, column=1).value == "AWS TOTAL"
        # Total tokens = 3 + 1 = 4
        assert ws.cell(row=total_row, column=8).value == 4
        wb.close()

    def test_summary_headers(self, tmp_path):
        """Summary headers match expected columns."""
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, [], {}, [], "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Summary"]
        headers = [cell.value for cell in ws[1]]
        expected = [
            "Account ID", "DDI Objects", "DDI Tokens", "Active IPs",
            "IP Tokens", "Managed Assets", "Asset Tokens", "Total Tokens",
        ]
        assert headers == expected
        wb.close()


class TestXlsxWarningsSheet:
    """Test the Warnings sheet for error handling."""

    def test_warnings_sheet_exists(self, tmp_path):
        """Warnings sheet is present in the workbook."""
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, [], {}, [], "aws")
        wb = openpyxl.load_workbook(filepath)
        assert "Warnings" in wb.sheetnames
        wb.close()

    def test_warnings_empty_errors(self, tmp_path):
        """Warnings sheet shows 'no errors' message when no errors."""
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, [], {}, [], "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Warnings"]
        assert ws.cell(row=2, column=1).value == "No errors occurred during discovery"
        wb.close()

    def test_warnings_with_errors(self, tmp_path):
        """Warnings sheet lists each error as a row."""
        errors = _sample_errors()
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, [], {}, errors, "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Warnings"]
        assert ws.cell(row=2, column=1).value == "333333333333"
        assert ws.cell(row=2, column=2).value == "eu-west-1"
        assert ws.cell(row=2, column=3).value == "ec2-instance"
        assert "AccessDenied" in ws.cell(row=2, column=4).value
        assert ws.cell(row=2, column=5).value == "Check IAM role permissions"
        wb.close()

    def test_warnings_multiple_errors(self, tmp_path):
        """Multiple errors produce multiple rows."""
        errors = [
            {"account": "111", "region": "us-east-1", "resource_type": "vpc",
             "error": "Timeout", "suggestion": "Retry"},
            {"account": "222", "region": "eu-west-1", "resource_type": "subnet",
             "error": "Throttled", "suggestion": "Wait"},
        ]
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, [], {}, errors, "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Warnings"]
        row_count = 0
        for row in ws.iter_rows(min_row=2, max_col=1):
            if row[0].value is not None:
                row_count += 1
        assert row_count == 2
        wb.close()


class TestXlsxEdgeCases:
    """Test edge cases for XLS report generation."""

    def test_empty_resources_and_summaries(self, tmp_path):
        """Report generates cleanly with no resources or summaries."""
        filepath = str(tmp_path / "test_report.xlsx")
        result = write_xlsx_report(filepath, [], {}, [], "aws")
        assert os.path.exists(result)
        wb = openpyxl.load_workbook(filepath)
        assert "Detail" in wb.sheetnames
        assert "Summary" in wb.sheetnames
        assert "Warnings" in wb.sheetnames
        wb.close()

    def test_resource_with_no_ips(self, tmp_path):
        """Resources with no IPs show empty IP address and 0 count."""
        resources = [_make_resource(ip_addresses=[])]
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, resources, {}, [], "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Detail"]
        assert ws.cell(row=2, column=6).value in ("", None)  # IP Addresses
        assert ws.cell(row=2, column=7).value == 0   # IP Count
        wb.close()

    def test_resource_with_no_tags(self, tmp_path):
        """Resources with no tags show empty tags column."""
        resources = [_make_resource(tags={})]
        filepath = str(tmp_path / "test_report.xlsx")
        write_xlsx_report(filepath, resources, {}, [], "aws")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Detail"]
        assert ws.cell(row=2, column=11).value in ("", None)  # Tags
        wb.close()


# ---- Estimator CSV Tests ----


def _read_csv(filepath: str) -> list[list[str]]:
    """Read a CSV file and return all rows."""
    with open(filepath, "r") as f:
        return list(csv.reader(f))


class TestEstimatorCsvGeneration:
    """Test estimator CSV file creation and structure."""

    def test_csv_creates_file(self, tmp_path):
        """CSV file is created and non-empty."""
        filepath = str(tmp_path / "estimator.csv")
        result = write_estimator_csv(filepath, _sample_account_summaries(), "aws")
        assert result == filepath
        assert os.path.exists(filepath)
        assert os.path.getsize(filepath) > 0

    def test_csv_returns_filepath(self, tmp_path):
        """write_estimator_csv returns the filepath written."""
        filepath = str(tmp_path / "estimator.csv")
        result = write_estimator_csv(filepath, {}, "aws")
        assert result == filepath

    def test_csv_header_row(self, tmp_path):
        """CSV has correct header columns."""
        filepath = str(tmp_path / "estimator.csv")
        write_estimator_csv(filepath, _sample_account_summaries(), "aws")
        rows = _read_csv(filepath)
        expected = [
            "Account ID", "DDI Objects", "Active IPs", "Managed Assets",
            "DDI Tokens", "IP Tokens", "Asset Tokens", "Total Tokens",
        ]
        assert rows[0] == expected

    def test_csv_row_count(self, tmp_path):
        """CSV has header + one row per account + TOTAL row."""
        summaries = _sample_account_summaries()
        filepath = str(tmp_path / "estimator.csv")
        write_estimator_csv(filepath, summaries, "aws")
        rows = _read_csv(filepath)
        assert len(rows) == 1 + len(summaries) + 1  # header + accounts + total

    def test_csv_totals_match_account_sums(self, tmp_path):
        """TOTAL row sums match individual account values."""
        summaries = _sample_account_summaries()
        filepath = str(tmp_path / "estimator.csv")
        write_estimator_csv(filepath, summaries, "aws")
        rows = _read_csv(filepath)

        total_row = rows[-1]
        assert total_row[0] == "TOTAL"

        # Sum from account rows (skip header at index 0 and total at -1)
        account_rows = rows[1:-1]
        for col_idx in range(1, 8):
            expected_sum = sum(int(row[col_idx]) for row in account_rows)
            assert int(total_row[col_idx]) == expected_sum

    def test_csv_empty_summaries(self, tmp_path):
        """CSV handles empty account summaries gracefully."""
        filepath = str(tmp_path / "estimator.csv")
        write_estimator_csv(filepath, {}, "aws")
        rows = _read_csv(filepath)
        assert len(rows) == 2  # header + total
        total_row = rows[-1]
        assert total_row[0] == "TOTAL"
        assert all(val == "0" for val in total_row[1:])

    def test_csv_accounts_sorted(self, tmp_path):
        """Account rows are sorted by account ID."""
        summaries = _sample_account_summaries()
        filepath = str(tmp_path / "estimator.csv")
        write_estimator_csv(filepath, summaries, "aws")
        rows = _read_csv(filepath)
        account_ids = [row[0] for row in rows[1:-1]]
        assert account_ids == sorted(account_ids)


# ---- Proof Manifest Tests ----


def _sample_scan_metadata() -> dict:
    """Create sample scan metadata."""
    return {
        "scan_timestamp": "2026-02-23T10:00:00Z",
        "scan_duration_seconds": 120,
    }


class TestProofManifestGeneration:
    """Test proof manifest file creation and structure."""

    def test_manifest_creates_file(self, tmp_path):
        """Manifest file is created and non-empty."""
        filepath = str(tmp_path / "proof.json")
        result = write_proof_manifest(
            filepath, _sample_resources(), _sample_account_summaries(),
            _sample_scan_metadata(), "aws",
        )
        assert result == filepath
        assert os.path.exists(filepath)
        assert os.path.getsize(filepath) > 0

    def test_manifest_returns_filepath(self, tmp_path):
        """write_proof_manifest returns the filepath written."""
        filepath = str(tmp_path / "proof.json")
        result = write_proof_manifest(filepath, [], {}, {}, "aws")
        assert result == filepath

    def test_manifest_valid_json(self, tmp_path):
        """Manifest is valid JSON."""
        filepath = str(tmp_path / "proof.json")
        write_proof_manifest(
            filepath, _sample_resources(), _sample_account_summaries(),
            _sample_scan_metadata(), "aws",
        )
        with open(filepath) as f:
            manifest = json.load(f)
        assert isinstance(manifest, dict)

    def test_manifest_required_fields(self, tmp_path):
        """Manifest contains all required top-level fields."""
        filepath = str(tmp_path / "proof.json")
        write_proof_manifest(
            filepath, _sample_resources(), _sample_account_summaries(),
            _sample_scan_metadata(), "aws",
        )
        with open(filepath) as f:
            manifest = json.load(f)
        required_fields = [
            "version", "provider", "scan_timestamp", "scan_duration_seconds",
            "scope", "ratios", "counts", "tokens",
            "resource_hash", "manifest_hash",
        ]
        for field in required_fields:
            assert field in manifest, f"Missing field: {field}"

    def test_manifest_scope_fields(self, tmp_path):
        """Scope section contains accounts, regions, resource types."""
        filepath = str(tmp_path / "proof.json")
        write_proof_manifest(
            filepath, _sample_resources(), _sample_account_summaries(),
            _sample_scan_metadata(), "aws",
        )
        with open(filepath) as f:
            manifest = json.load(f)
        scope = manifest["scope"]
        assert scope["accounts_scanned"] == 2
        assert scope["regions_scanned"] == 2
        assert isinstance(scope["resource_types"], list)
        assert len(scope["resource_types"]) > 0

    def test_manifest_ratios(self, tmp_path):
        """Ratios match hardcoded token constants."""
        filepath = str(tmp_path / "proof.json")
        write_proof_manifest(filepath, [], {}, {}, "aws")
        with open(filepath) as f:
            manifest = json.load(f)
        assert manifest["ratios"]["ddi_per_token"] == 25
        assert manifest["ratios"]["ips_per_token"] == 13
        assert manifest["ratios"]["assets_per_token"] == 3

    def test_manifest_counts(self, tmp_path):
        """Counts section has correct resource totals."""
        resources = _sample_resources()
        filepath = str(tmp_path / "proof.json")
        write_proof_manifest(
            filepath, resources, _sample_account_summaries(),
            _sample_scan_metadata(), "aws",
        )
        with open(filepath) as f:
            manifest = json.load(f)
        counts = manifest["counts"]
        assert counts["total_resources"] == len(resources)
        counted = sum(1 for r in resources if r.counted)
        assert counts["counted_resources"] == counted
        assert counts["skipped_resources"] == len(resources) - counted

    def test_manifest_resource_hash_reproducible(self, tmp_path):
        """Same input produces same resource_hash (deterministic)."""
        resources = _sample_resources()
        summaries = _sample_account_summaries()
        metadata = _sample_scan_metadata()

        filepath1 = str(tmp_path / "proof1.json")
        filepath2 = str(tmp_path / "proof2.json")
        write_proof_manifest(filepath1, resources, summaries, metadata, "aws")
        write_proof_manifest(filepath2, resources, summaries, metadata, "aws")

        with open(filepath1) as f:
            m1 = json.load(f)
        with open(filepath2) as f:
            m2 = json.load(f)
        assert m1["resource_hash"] == m2["resource_hash"]

    def test_manifest_hash_changes_on_different_input(self, tmp_path):
        """manifest_hash changes when resource data changes."""
        resources1 = _sample_resources()
        resources2 = _sample_resources()[:3]  # Fewer resources

        filepath1 = str(tmp_path / "proof1.json")
        filepath2 = str(tmp_path / "proof2.json")
        write_proof_manifest(
            filepath1, resources1, _sample_account_summaries(),
            _sample_scan_metadata(), "aws",
        )
        write_proof_manifest(
            filepath2, resources2, _sample_account_summaries(),
            _sample_scan_metadata(), "aws",
        )

        with open(filepath1) as f:
            m1 = json.load(f)
        with open(filepath2) as f:
            m2 = json.load(f)
        assert m1["manifest_hash"] != m2["manifest_hash"]

    def test_manifest_hash_is_valid_sha256(self, tmp_path):
        """resource_hash and manifest_hash are valid SHA-256 hex strings."""
        filepath = str(tmp_path / "proof.json")
        write_proof_manifest(
            filepath, _sample_resources(), _sample_account_summaries(),
            _sample_scan_metadata(), "aws",
        )
        with open(filepath) as f:
            manifest = json.load(f)
        # SHA-256 hex digest is 64 characters
        assert len(manifest["resource_hash"]) == 64
        assert len(manifest["manifest_hash"]) == 64
        # Verify all hex chars
        int(manifest["resource_hash"], 16)
        int(manifest["manifest_hash"], 16)

    def test_manifest_hash_verifiable(self, tmp_path):
        """manifest_hash can be verified by recomputing from manifest content."""
        filepath = str(tmp_path / "proof.json")
        write_proof_manifest(
            filepath, _sample_resources(), _sample_account_summaries(),
            _sample_scan_metadata(), "aws",
        )
        with open(filepath) as f:
            manifest = json.load(f)

        saved_hash = manifest.pop("manifest_hash")
        recomputed = hashlib.sha256(
            json.dumps(manifest, sort_keys=True, default=str).encode()
        ).hexdigest()
        assert recomputed == saved_hash

    def test_manifest_empty_resources(self, tmp_path):
        """Manifest handles empty resource list gracefully."""
        filepath = str(tmp_path / "proof.json")
        write_proof_manifest(filepath, [], {}, _sample_scan_metadata(), "aws")
        with open(filepath) as f:
            manifest = json.load(f)
        assert manifest["counts"]["total_resources"] == 0
        assert manifest["counts"]["counted_resources"] == 0
        assert manifest["counts"]["skipped_resources"] == 0
        assert manifest["tokens"]["total_tokens"] == 0

    def test_manifest_version(self, tmp_path):
        """Manifest version is 1.0."""
        filepath = str(tmp_path / "proof.json")
        write_proof_manifest(filepath, [], {}, {}, "aws")
        with open(filepath) as f:
            manifest = json.load(f)
        assert manifest["version"] == "1.0"

    def test_manifest_provider(self, tmp_path):
        """Manifest records the provider name."""
        filepath = str(tmp_path / "proof.json")
        write_proof_manifest(filepath, [], {}, {}, "aws")
        with open(filepath) as f:
            manifest = json.load(f)
        assert manifest["provider"] == "aws"
